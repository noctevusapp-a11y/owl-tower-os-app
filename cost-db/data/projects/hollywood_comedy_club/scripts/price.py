#!/usr/bin/env python3
"""Hollywood Comedy Club TI - price the takeoff from the company's book-derived cost table.

Rules applied (owner's standing rules):
  1. a line is priced ONLY from rows with status = verified; the Los Angeles latest-month columns are used
     (material_la_latest, labor_la_latest, equipment_la_latest, total_la_latest, total_incl_op_la_latest)
  2. a machine_read / flagged row is never used as a price: it is reported as candidate_row with the price empty
  3. no invented numbers, no allowances, no tolerances; every unit conversion is written out (factor column)
  4. no second location or time factor: the *_la_latest columns already carry the LA index and BLS escalation
  5. bare components are kept separate; the book O&P column is reported for reference only

Price tiers (price_status):
  priced_exact          the verified row describes the specified work AND the product / size / type as the issued
                        sheets fix it, same unit or an exact stated conversion
  priced_substitute     same kind of work, but the product, size, gauge, type or finish system is not fixed by the
                        issued sheets or differs from the row (match = substitute), or the row covers part of the
                        line only (match = partial); the exact difference is written in `reason`
  book_row_not_matching the nearest book row is a different kind of work or product: unpriced, candidate_row filled
  cost_code_duplicated  the cost code is carried by more than one verified row in the table: unpriced
  not_quantified / spec_not_issued / not_in_book / book_row_not_verified / owner_provided / quote / none_required
Cost codes: the book is loaded with ALL rows per cost_code; a code is used only if exactly one row carries it with
status = verified and that row's own description matches the pattern recorded for the component.

Run:  python3 -I price.py [path/to/book_lines_la_latest.csv]
Writes ../estimate_la_latest.csv, ../estimate_la_latest.xlsx, ../unpriced_items.csv, ../_totals.json
"""
import csv, json, os, re, sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
sys.path.insert(0, HERE)                 # explicit: python -I does not add the script directory
import takeoff as TO

BOOK = sys.argv[1] if len(sys.argv) > 1 else "/home/user/owl-tower-os-app/cost-db/data/books/rsmeans_bcc_2022/book_lines_la_latest.csv"
OUT_CSV = os.path.join(PROJ, "estimate_la_latest.csv")
OUT_XLSX = os.path.join(PROJ, "estimate_la_latest.xlsx")
OUT_UNP = os.path.join(PROJ, "unpriced_items.csv")
OUT_TOT = os.path.join(PROJ, "_totals.json")

# ---------------------------------------------------------------------------------------------
# Book: every row per cost code (no first-occurrence-wins)
# ---------------------------------------------------------------------------------------------
BOOKROWS = defaultdict(list)
with open(BOOK, newline="") as f:
    for r in csv.DictReader(f):
        if r["cost_code"]:
            BOOKROWS[r["cost_code"]].append(r)


def D(x):
    return Decimal(x) if x not in ("", None) else None


def q2(x):
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------------------------------------
# Pricing map
#   C(cost_code, factor, component description, tier, reason, expect[, match])   priced component
#       factor converts the takeoff quantity into the book quantity (book_qty = takeoff qty x factor)
#       tier = "exact" | "substitute"; match defaults to the tier name, "partial" for part coverage
#       expect = regex that the verified row's own description must match (guards against duplicate codes)
#   U(component description, status, reason, candidate)                           unpriced component
# ---------------------------------------------------------------------------------------------
P = {}


def C(cc, factor, desc, tier, reason, expect, match=None):
    return dict(kind="comp", cc=cc, factor=factor, desc=desc, tier=tier, match=(match or tier), reason=reason, expect=expect)


def U(desc, status, reason, candidate=""):
    return dict(kind="ucomp", desc=desc, status=status, reason=reason, candidate=candidate)


def priced(item_id, *comps):
    P[item_id] = dict(kind="priced", comps=list(comps))


def unpriced(item_id, status, reason, candidate=""):
    P[item_id] = dict(kind="unpriced", status=status, reason=reason, candidate=candidate)


GWB_X = ("OT-09291030-2150", "5/8\" fire-resistant (Type X) GWB on walls, taped and finished level 4 (S.F.)", r"level 4")
INSUL = ("OT-07211620-1320", "Mineral wool batts 3-1/2\" thick R15, wall or ceiling (S.F.)", r"3-1/2\" thick, R15")

# ---- Demolition
unpriced("D01", "not_in_book", "equipment count per K sheets not issued; the only hood-removal row is flagged", "OT-11050510-4750 (flagged)")
priced("D02", C("OT-06050520-1500", 1, "Counter top removal, straight runs (L.F.)", "substitute",
                "drawn: bar, diewall, back bar and millwork removal; the row covers the counter top only", r"Counter top, straight runs", match="partial"))
unpriced("D03", "not_quantified", "not dimensioned on A-1.10")
unpriced("D04", "not_in_book", "no stair demolition row in the book")
unpriced("D05", "not_quantified", "wall SF not dimensioned on A-1.10 / A-1.20")
unpriced("D06", "not_quantified", "wall LF not dimensioned on A-1.10 / A-1.20")
unpriced("D07", "book_row_not_verified", "coffer material not stated; GWB ceiling demolition row is flagged; plaster-on-lath row (OT-09050510-1000, verified) would apply only if the coffers are plaster", "OT-09050510-0200 (flagged)")
unpriced("D08", "not_quantified", "trench and opening dimensions per P sheets not issued", "OT-02411916-0020 (verified, C.F.)")
unpriced("D09", "not_in_book", "no floor-opening patch row; opening size not stated")
unpriced("D10", "book_row_not_verified", "aluminum window removal rows are flagged (p268 'Aluminum, including trim, to 12 S.F.' sits under an unreadable section header and carries no cost code)", "p268 window demolition line 0200 (flagged, no cost code)")
unpriced("D11", "book_row_not_verified", "interior door removal rows are machine_read / flagged", "OT-08050510-1020 (machine_read)")
unpriced("D12", "not_quantified", "fixture count not shown on A-1.20")
unpriced("D13", "not_quantified", "not quantifiable")
unpriced("D14", "not_quantified", "not quantifiable")
unpriced("D15", "not_quantified", "extent not shown")
unpriced("D16", "not_quantified", "count not shown")

# ---- Interior walls (components priced on the one-face wall SF; GWB both sides = factor 2)
priced("W01",
       C("OT-09221613-3200", 1, "Metal studs, non-load bearing, galv., 12' high, 20 ga., 3-5/8\" wide, 16\" OC (S.F.)", "substitute",
         "drawn types 1 and 2 are 3-5/8\" 20 ga 16\" OC as the row; the class A LF also contains AC walls (4\" 16 ga) and type 5 (2\" furring on (E) wall) that are not separable; row stud height 12' vs 13.5 ft class average", r"3-5/8\" wide, 16\" OC"),
       C(GWB_X[0], 2, GWB_X[1] + ", both sides", "substitute",
         "drawn: 5/8\" Type X each side (types 1, 2) as the row, but the finish is drawn Level 5 in public areas (row: level 4) and the AC (2 layers each side) / type 5 (one side) layer counts are not separable from the class A LF", GWB_X[2]),
       C(INSUL[0], 1, INSUL[1] + " as the 4\" acoustic insulation", "substitute",
         "drawn: '4\" acoustic insulation' with no product; row is 3-1/2\" R15 mineral wool batt", INSUL[2]))
priced("W02",
       U("Metal studs 6\" 20 ga 16\" OC (type 1A / type 3)", "book_row_not_matching",
         "drawn: 6\" 20 ga studs 16\" OC; nearest verified row OT-09221613-3410 is 25 ga 6\" studs 12\" OC under the 16' high 25 ga group; the 20 ga 6\" rows carry an unreadable unit", "OT-09221613-3410 (verified, 25 ga 12\" OC); OT-09221613-3490 (verified, unit unreadable)"),
       C(GWB_X[0], 4, GWB_X[1] + ", 2 layers each side", "substitute",
         "drawn (type 1A): '2 layers TYPE-X GWB on ea. side' with no thickness stated; base layers priced with the taped-and-finished row (no verified unfinished Type X row); type 3 exterior plaster side not separable (W05)", GWB_X[2]))
unpriced("W03", "spec_not_issued", "shaft/enclosure assembly per A-6.xx / S sheets not issued")
unpriced("W04", "not_quantified", "AC-tagged LF not separable from the class A total")
unpriced("W05", "not_quantified", "type 3 LF not separable from the class B total", "OT-09242340-1000 (stucco 3 coats, verified, S.Y.); OT-09242340-1900 (mesh wired to steel, verified)")
unpriced("W06", "not_quantified", "not dimensioned")
unpriced("W07", "not_quantified", "not dimensioned")
priced("W08",
       U("Metal studs 6\" 16 ga 16\" OC (exit stair 1-hr wall)", "book_row_not_matching",
         "drawn: 6\" 16 ga studs 16\" OC; nearest verified S.F. row OT-09221613-3410 is 25 ga 6\" 12\" OC; the 16 ga rows in 05 41 13 are priced per stud (Ea.)", "OT-09221613-3410 (verified, 25 ga 12\" OC); OT-05411310-0330 (verified, 16 ga 6\" wide, Ea.)"),
       C(GWB_X[0], 2, GWB_X[1] + ", both sides", "exact", "drawn: '5/8\" TYPE-X GWB ea. side' (submittal wall note 1)", GWB_X[2]),
       C(INSUL[0], 1, INSUL[1] + " as the 4\" acoustic insulation", "substitute",
         "drawn: '4\" acoustic insulation, UL Des U404' with no product; row is 3-1/2\" R15 mineral wool batt", INSUL[2]))
unpriced("W09", "not_quantified", "public / back-of-house wall split not separable")

# ---- Ceilings
priced("C01",
       U("New steel ceiling joists 3-5/8\" (span under 4') / 6\" with track", "book_row_not_matching",
         "drawn: GWB on new steel joists; nearest verified row OT-09222613-3400 is suspended 1-5/8\" carrier channels, a different framing system", "OT-09222613-3400 (verified, suspended channels)"),
       C("OT-09291030-3090", 1, "5/8\" GWB on ceilings, standard, with compound skim coat level 5 (S.F.)", "substitute",
         "drawn: 5/8\" Ultralight Firecode 30 GWB, smooth washable finish; row is 5/8\" standard board (the fire-resistant ceiling rows are machine_read); level 5 taken as the smooth finish", r"level 5"))
priced("C02", C("OT-09912372-1670", 1, "Painting, 2 coats finish on drywall or plaster with roller, incl. surface prep and primer (S.F.)", "substitute",
                "P-1 paint system per A-7.20 not issued; walls row applied to ceilings", r"2 coats finish, on drywall or plaster"))
unpriced("C03", "not_quantified", "soffits not dimensioned")
priced("C04", C("OT-09912372-1670", 1, "Painting, 2 coats finish on drywall or plaster with roller (S.F.)", "substitute",
                "refinish of (E) GWB ceilings; finish per A-7.20 not issued", r"2 coats finish, on drywall or plaster"))
priced("C05", C("OT-03051320-1660", Decimal(1) / Decimal(350), "Concrete sealer, water based, 350 S.F./gal., 5 gallon pail (Gal.) - material only", "substitute",
                "coverage printed on the row: description_path 'Sealer, water based, 350 S.F, 55 gallon drum | 5 gallon pail' (350 S.F. per gallon); factor = 1/350 gal per S.F.; stain and sealer products ST-2 / SL-2 per A-7.20 not issued; cleaning, patching and application labor have no verified row",
                r"5 gallon pail", match="partial"))
unpriced("C06", "not_in_book", "no K-13 / spray cellulose acoustic insulation row in the book")
unpriced("C07", "not_quantified", "hatch count not shown", "OT-08311310-4150 (verified, recessed drywall door 12x24)")
unpriced("C08", "not_quantified", "not dimensioned")

# ---- Flooring
unpriced("F01", "spec_not_issued", "floor finish code not issued (A-2.75 / A-7.20)")
unpriced("F02", "spec_not_issued", "floor finish code not issued (A-2.75 / A-7.20)")
unpriced("F03", "spec_not_issued", "lobby tile (FL-7/FL-8) and green room FL-4 areas not separable; tile product per A-7.20 not issued", "OT-09311310-3300 (verified porcelain tile)")
unpriced("F04", "spec_not_issued", "kitchen floor finish not issued (quarry tile / epoxy TBD in prior estimate)", "OT-09321610-0900 (quarry tile 6x6, verified); OT-09672323-1500 (epoxy topping 1/4\", verified)")
unpriced("F05", "spec_not_issued", "floor finish code not issued (A-2.76 / A-7.20)")
priced("F06", C("OT-09311310-3300", 1, "Ceramic tile, porcelain type, thin set, 1 color, color group 2, 1\" x 1\" (S.F.)", "substitute",
                "WT-1 tile product and size per A-7.20 not issued; row is 1\" x 1\" porcelain mosaic", r"porcelain type"))
unpriced("F07", "not_in_book", "custom millwork platform with LVT; no verified platform or LVT row")
unpriced("F08", "not_quantified", "not dimensioned")
priced("F09", C("OT-06221315-5105", 1, "Base molding, classic profile 5/8\" x 5-1/2\", finger jointed and primed, poplar (L.F.)", "substitute",
                "B-4 base profile and species per A-7.20 not issued; paint P-2 on the base not priced (no verified trim-paint row)", r"Poplar"))
unpriced("F10", "not_in_book", "LED base profile not in the book")
priced("F11", C("OT-09311310-0700", 1, "Thin-set ceramic tile cove base 4-1/4\" x 4-1/4\" (L.F.)", "substitute",
                "B-1 tile base product per A-7.20 not issued", r"Cove base, 4-1/4"))
unpriced("F12", "not_quantified", "no finish plan issued")
unpriced("F13", "not_quantified", "length not dimensioned")

# ---- New roofing
unpriced("R01", "not_quantified", "no roof plan or M sheets issued")

# ---- Interior finishes
priced("I01", C("OT-09721613-3300", 1, "Vinyl wall covering, fabric-backed, medium weight, type 2 (20-24 oz./S.Y.) (S.F.)", "substitute",
                "the issued sheets tag WC-8 only; 'Type II' comes from the prior session's note (mac_session_page3.txt L220), A-7.20 not issued; book line notes 'add 10-30% waste at takeoff', no waste added", r"Medium weight, type 2"))
unpriced("I02", "book_row_not_verified",
         "drawn: 1/2\" x 1/4\" poplar strip stained to match WC-8; the only poplar crown row (p203 line 2510, 'Crown, 5/8\" x 5/8\", poplar', formerly OT-06221330-2510) is flagged in the replaced table (SECTION_ID_DOUBTFUL, no cost code, no total); the verified poplar rows of 06 22 13.30 (lines 0455-0575) are casing profiles 3/4\" x 2-1/2\" to 9-1/2\", a different product",
         "p203 line 2510 (flagged, no cost code); OT-06221330-0575 (verified, casing 3/4\" x 2-1/2\" poplar, not matching)")
priced("I03", C("OT-09722310-3700", 1, "Wallpaper, average workmanship, solid pattern, low cost paper (S.F.)", "substitute",
                "WC-3 product per A-7.20 not issued; row is a low-cost paper", r"Wallpaper, average workmanship"))
priced("I04", C("OT-09722310-3700", 1, "Wallpaper, average workmanship, solid pattern, low cost paper (S.F.)", "substitute",
                "WC-4 'WALL PAPER' per Addendum; product per A-7.20 not issued; row is a low-cost paper", r"Wallpaper, average workmanship"))
priced("I05", U("TX-4 fabric-wrapped 2\" acoustic panels", "book_row_not_matching",
                "drawn: 2\" acoustic panels with orange fabric stretched over them (A-4.22A); nearest verified row OT-09841310-1000 sits under 'FIXED PANELS, perforated steel facing' (mineral fiberboard, fabric covered, 3/4\" thick, concealed spline), a different product", "OT-09841310-1000 (verified)"))
unpriced("I06", "not_in_book", "drapery fabric not in the book (hardware rods only)", "OT-12221610-4040 (traverse rod, verified)")
unpriced("I07", "not_in_book", "no verified wood paneling row (hardboard paneling only)")
priced("I08", C("OT-09912372-1670", 1, "Painting walls, 2 coats on drywall with roller incl. primer (S.F.)", "substitute",
                "P-1 paint system per A-7.20 not issued", r"2 coats finish, on drywall or plaster"))
priced("I09", C("OT-09311310-3300", 1, "Ceramic tile, porcelain type, thin set, 1 color, 1\" x 1\" (S.F.) on walls", "substitute",
                "WT-1 product and size per A-7.20 not issued; floor-tile row applied to the wainscot", r"porcelain type"))
priced("I10", C("OT-09912333-2200", 1, "Flush door and frame 3' x 7', oil, primer, brushwork, paint 1 coat (Ea.)", "substitute",
                "book row is one finish coat on a flush door and frame; P-5 system and door type per A-7.20 / A-7.10 not issued", r"Paint, 1 coat", match="partial"))
unpriced("I11", "not_quantified", "moulding LF per door not dimensioned (A-7.10 not issued)")
unpriced("I12", "not_in_book", "no verified casing / cased-opening row")
priced("I13", C("OT-09773010-0030", 1, "Fiberglass reinforced plastic panels .090\" on walls, adhesive mounted, smooth surface (S.F.)", "substitute",
                "drawn: 'PROVIDE FRP' with no thickness, surface or fastening stated", r"Smooth surface"))
unpriced("I14", "not_quantified", "column surface area not stated; sealer material carried in C05")
unpriced("I15", "not_in_book", "custom perforated metal enclosure")
unpriced("I16", "not_quantified", "custom millwork, not dimensioned")
unpriced("I17", "not_quantified", "not dimensioned")
priced("I18", C("OT-09912372-1670", 1, "Painting, 2 coats on drywall or plaster with roller (S.F.)", "substitute",
                "P-3 on 2\" painted wood frames; drywall/plaster walls row applied to wood; system per A-7.20 not issued", r"2 coats finish, on drywall or plaster"))

# ---- Exterior finishes
priced("X01",
       C("OT-08431310-0020", 1, "Aluminum-framed entrance door, 3' x 7' opening, clear anodized finish (Opng.)", "substitute",
         "drawn: tempered glass storefront exit door to match (E) entry doors in appearance, size and configuration; finish and size per A-7.10 not issued; row is a stock clear-anodized aluminum entrance door", r"clear anodized"),
       C("OT-08431310-0620", 1, "Door frame for above doors, 3'-0\" x 7'-0\", bronze finish (Opng.)", "substitute",
         "frame finish to match existing not fixed; row is bronze finish", r"Bronze finish"),
       U("Panic exit device with accessible lever trim", "book_row_not_matching",
         "drawn: panic hardware and accessible hardware on exterior exit doors (submittal KN3, A-4.01 KN5); the only verified rows with a panic device are average hardware sets for school / motel doors (OT-08712015-2250, 'Door'), a different product", "OT-08712015-2250 (verified, hardware set)"),
       C("OT-08712030-8030", 1, "Door closer, surface mounted, light duty, parallel arm, primed (Ea.)", "substitute",
         "drawn: self-closing device (submittal KN3); closer grade not fixed; row is light duty", r"Light duty"),
       C("OT-08712065-0600", 1, "Bronze panic threshold 5\" wide, 1/2\" thick (Ea.)", "substitute",
         "drawn: 'exterior weatherstripping and threshold at bottom of door'; threshold type not fixed; weatherstripping not priced (no verified row)", r"panic threshold"))
unpriced("X02", "not_quantified", "panel sizes not dimensioned", "OT-08811010-0800 (1/4\" tempered glass, verified); OT-08411320-0150 (tube framing, verified)")
unpriced("X03", "not_quantified", "not dimensioned")
unpriced("X04", "not_quantified", "not dimensioned", "OT-09242340-1000 (stucco, verified)")
unpriced("X05", "not_quantified", "not dimensioned")
unpriced("X06", "not_quantified", "not dimensioned", "OT-05521350-0915 (wall rail, verified)")

# ---- Doors & hardware
priced("DR1",
       C("OT-08131315-0080", 1, "Steel door, flush, \"B\" label 90 min, full panel, 20 ga., 3'-0\" x 7'-0\" (Ea.)", "substitute",
         "door types, sizes and ratings per A-7.10 not issued (prior estimate: HM rated + wood); the verified wood flush-door rows carry no unit and the 08 14 16 rows are flagged; row is a 3'-0\" x 7'-0\" steel B-label door", r"3'-0\" x7"),
       C("OT-08121313-3600", 1, "Standard hollow metal frame, 14 ga., up to 5-3/4\" jamb depth, 4'-0\" x 7'-0\" single (Ea.)", "substitute",
         "size difference: row is a 4'-0\" x 7'-0\" single frame, doors are drawn as 3'-0\" leaves (A-2.10 '3'-0\" DOOR'); no verified 3'-0\" x 7'-0\" single frame row (OT-08121313-0100 is machine_read); gauge per A-7.10 not issued", r"4'-0\" x 7'-0\" single"),
       C("OT-08712040-0500", 1, "Lockset, standard duty, lever handled, keyed, single cylinder function (Ea.)", "substitute",
         "lock functions per A-7.10 not issued; one keyed lever lockset per door", r"Lever handled, keyed"),
       U("Hinges, closers and remaining hardware per A-7.10", "book_row_not_matching",
         "hardware sets per A-7.10 not issued; no verified hinge row; OT-08712030-8030 is a surface closer under the 'electronic closer-holder' heading and is not a hinge set", "OT-08712030-8030 (verified, closer)"))
priced("DR2",
       C("OT-08131315-0080", 1, "Steel door, flush, \"B\" label 90 min, full panel, 20 ga., 3'-0\" x 7'-0\" (Ea.)", "substitute",
         "drawn: 1-hr fire-rated exit access door, size per A-7.10 not issued; row is a 90-minute (B label) 3'-0\" x 7'-0\" steel door", r"3'-0\" x7"),
       C("OT-08121313-5400", 1, "Hollow metal frame 14 ga., \"B\" label, up to 5-3/4\" deep, 4'-0\" x 7'-0\" single (Ea.)", "substitute",
         "size difference: row is a 4'-0\" wide labeled frame; door size per A-7.10 not issued", r"\"B\" label"),
       U("Panic hardware on the rated exit access door", "book_row_not_matching",
         "drawn: panic hardware (submittal KN5); the only verified panic rows are average hardware sets for school / motel doors (OT-08712015-2250), a different product", "OT-08712015-2250 (verified, hardware set)"))
unpriced("DR3", "not_quantified", "count per mech/plumb drawings not issued", "OT-08311310-4150 (verified)")
priced("DR4", C("OT-08912610-0050", 1, "Door louver 12\" x 12\" with powder coat (Ea.)", "substitute", "louver size and finish not given on A-4.30", r"12\"x 12\", with powder coat"))

# ---- Stairs & railing
priced("S01", C("OT-05511350-0200", 1, "Pan stairs, shop fabricated, steel stringers, metal pan tread for concrete infill, picket rail, 3'-6\" wide (Riser)", "substitute",
                "drawn: 3'-0\" wide stair with C12x20.7 stringers and 1-1/2\" handrails; row is 3'-6\" wide with a picket rail", r"Metal pan tread for concrete infill"))
unpriced("S02", "not_in_book", "no verified stair-pan concrete fill row")
priced("S03", C("OT-05122317-5600", Decimal(1) / Decimal(12), "Structural tubing, rect., 8\" x 4\" x 3/8\" x 12'-0\" (Ea. = one 12'-0\" piece)", "substitute",
                "piece length printed on the row: '8\" x 4\" x 3/8\" x 12'-0\"'; factor 1/12 per LF (18 LF = 1.5 pieces); wall thickness differs: drawn 1/4\", row 3/8\"", r"8\" x 4\" x 3/8\" x 12'-0"))
unpriced("S04", "not_in_book", "no verified HSS 4x4 row")
unpriced("S05", "not_in_book", "landing joists are listed per stud (Ea.) only; Sure-Board decking not in the book")
priced("S06", U("Embedded anchors 1/2\" x 4\" and threaded rods at base plates", "book_row_not_matching",
                "drawn: 1/2\" x 4\" embedded anchors / threaded rods (S-3.10); nearest verified row OT-05051920-0600 is a 1/2\" x 2-1/16\" zinc expansion shield, a different product", "OT-05051920-0600 (verified, expansion shield)"))
priced("S07", C("OT-05521350-0915", 1, "Wall rail, aluminum pipe 1-1/2\" diameter, satin finish (L.F.)", "substitute",
                "drawn: 1-1/2\" diameter metal handrail at 34-38\" (material not stated); row is aluminum satin wall rail; one side may be carried by the picket rail in S01", r"1-1/2\" diameter, satin finish"))
unpriced("S08", "not_in_book", "visual warning strip not in the book")
unpriced("S09", "not_quantified", "not dimensioned")
unpriced("S10", "not_quantified", "A-6.10/11/12 of the main set not issued; riser count unknown")
unpriced("S11", "not_quantified", "LF not dimensioned; A-4.15 not issued", "OT-05521350-2050 (42\" steel rail, verified)")
unpriced("S12", "not_quantified", "not dimensioned")
unpriced("S13", "spec_not_issued", "post sizes per S sheets not issued")
unpriced("S14", "not_quantified", "not dimensioned")

# ---- Owner-provided equipment install
for k in ("O01", "O02", "O03", "O04", "O05", "O06"):
    unpriced(k, "owner_provided", "owner-provided item; setting/installation labor has no verified row")
unpriced("O07", "spec_not_issued", "window type per A-7.10 not issued")
unpriced("O08", "not_quantified", "length not dimensioned")

# ---- Lighting
for k in list(TO.LIGHTS_L1) + list(TO.LIGHTS_M):
    unpriced(f"L_{k}", "spec_not_issued", "fixture type per A-3.70 not issued")
unpriced("L_G", "spec_not_issued", "decorative fixtures per A-3.70 not issued")
unpriced("L_EX", "not_quantified", "symbols not countable; E plans not issued", "OT-26521316-0200 (LED exit sign, verified); OT-26521310-0900 (emergency unit, verified)")
unpriced("L_SW", "not_quantified", "symbols not countable", "OT-26272620-0200 (switch), OT-26272620-1650 (dimmer), verified")
priced("L_RC", C("OT-26272620-2460", 1, "Receptacle, duplex, 120 V, grounded, 15 A (Ea.) - device only", "substitute",
                 "device only; box, wiring and circuit not priced (E plans not issued); 'OUTLET, BLACK' colour/grade not fixed", r"Receptacle, duplex", match="partial"))
unpriced("L_LED", "not_in_book", "LED tape / base strip not in the book")
unpriced("L_BR", "not_quantified", "E plans not issued")

# ---- Mechanical
unpriced("M01", "book_row_not_verified", "water-source heat pump rows are machine_read; tonnage per M schedule not issued", "OT-23814610-2100 / -2140 / -2220 (machine_read)")
unpriced("M02", "not_in_book", "Type I kitchen hood not in the book (removal row only, flagged)", "OT-11050510-4750 (flagged, removal)")
unpriced("M03", "not_quantified", "duct LF not drawn; fire wrap not in the book")
unpriced("M04", "spec_not_issued", "capacities not available", "OT-23733910-0100..0300 (make-up air by MBH, verified)")
priced("M05", C("OT-23371360-1120", 1, "Register, air supply, one or two way deflection, adjustable curved face bars, 12\" x 12\" (Ea.)", "substitute",
                "diffuser / grille types and sizes per M-1 not on hand; all 45 priced as 12\" x 12\" supply registers", r"12\" x 12"))
unpriced("M06", "not_quantified", "M sheets not issued")
unpriced("M07", "not_quantified", "M sheets not issued")
priced("M08", C("OT-23342310-6680", 1, "Bath exhaust fan, light combination, squirrel cage, 100 W, 70 CFM (Ea.)", "substitute",
                "drawn: mechanical exhaust diffuser switched with the lights, capacity per M sheets not issued; row is a residential light-combination fan", r"Light combination"))
unpriced("M09", "not_quantified", "M sheets not issued")

# ---- Plumbing
priced("P01",
       C("OT-22411313-1140", 1, "Water closet, tank type, vitreous china, floor mounted, two piece, 1.28 gpf, ADA (Ea.)", "substitute",
         "drawn: accessible water closet (A-5.12 fixture schedule not readable in the text layer); flush type and model not fixed; row is a floor-mounted tank type", r"1.28 gpf, ADA"),
       C("OT-22411313-1980", 1, "Water closet rough-in, supply, waste and vent (Ea.)", "exact", "rough-in for one water closet as drawn", r"rough-in, supply, waste and vent"))
priced("P02",
       U("Lavatory K-2035-8 Pinoir, wall-hung vitreous china 20\" x 18\", with metering faucet", "book_row_not_matching",
         "drawn: wall-hung lavatory K-2035-8 Pinoir (A-2.20, A-5.12 KN8, KN9); nearest verified row OT-22441613-2900 is a vanity-top vitreous china 20\" x 16\" bowl, a different mounting; the wall-hung rows are not verified", "OT-22441613-2900 (verified, vanity top)"),
       C("OT-22441613-3580", 1, "Lavatory rough-in, supply, waste and vent (Ea.)", "exact", "rough-in for one lavatory as drawn", r"supply, waste and vent"))
unpriced("P03", "spec_not_issued", "interceptor capacity (GPM) per P sheets not issued; in the replaced table the grease interceptor rows (p514 lines 0150 / 0200 / 1060 / 1120, by GPM) are flagged under an unreadable section header and carry no cost code", "p514 lines 0150 / 0200 / 1060 / 1120 (flagged, no cost code)")
unpriced("P04", "not_quantified", "LF per P sheets not issued")
unpriced("P05", "not_quantified", "count per P sheets not issued")
unpriced("P06", "not_quantified", "K / P sheets not issued")
unpriced("P07", "not_quantified", "P sheets not issued")

# ---- Low voltage
unpriced("V01", "not_quantified", "no FA drawings issued")
unpriced("V02", "not_quantified", "no LV drawings issued")
unpriced("V03", "not_quantified", "head count not countable; F sheets not issued", "OT-21131350-3730 (sprinkler head, verified)")

# ---- F&B equipment
unpriced("E01", "not_in_book", "quote required (K sheets not issued)")
unpriced("E02", "not_in_book", "quote required")
unpriced("E03", "not_in_book", "no verified stainless steel countertop row")
unpriced("E04", "not_quantified", "A-5.20 not issued")

# ---- Other required scope
unpriced("G01", "book_row_not_verified", "field personnel rows are flagged", "OT-01311320-0240 (superintendent, header/flagged)")
unpriced("G02", "book_row_not_verified", "field personnel rows are flagged", "OT-01311320-0180 (project manager, flagged)")
unpriced("G03", "not_in_book", "no coordinator row")
for k in ("G04", "G05", "G06", "G10", "G11", "G12", "G13", "G14"):
    unpriced(k, "not_in_book", "company lump-sum item; no matching verified row (temporary utility rows carry no cost)")
priced("G07", C("OT-02411919-0800", 1, "Dumpster, weekly rental, 1 dump/week, 30 C.Y. capacity, 7 tons (Week)", "substitute",
                "dumpster size and dumps per week not stated in the prior estimate line; 30 C.Y. row used", r"30 CY"))
priced("G08", C("OT-01741320-0050", Decimal(TO.GROSS_SF) / Decimal(1000), "Cleanup of floor area, continuous, per day, during construction (M.S.F.) - 7.141 M.S.F. per day", "exact",
                "book qty = 150 days x 7.141 M.S.F. (7,141 SF gross / 1000); the row is continuous daily cleanup of the floor area, the work carried by the prior line 01-080", r"Cleanup of floor area, continuous"))
unpriced("G09", "book_row_not_verified", "final clean by GC row is flagged", "OT-01741320-0100 (flagged)")
unpriced("G15", "book_row_not_verified", "electric dumbwaiter rows are flagged", "OT-14121010-0020 (flagged)")
unpriced("G16", "quote", "vendor quote origin (McKinley Equipment e-mail, Garaventa ROM on file); not a book price")
for k in ("G17", "G18", "G19", "G20", "G21"):
    unpriced(k, "not_quantified", "drawings not issued")
priced("G22",
       C("OT-10441353-1100", 1, "Fire equipment cabinet, portable extinguisher, single, 8\" x 12\" x 27\", steel door and frame (Ea.)", "substitute",
         "cabinet type (recessed/surface, door style) not specified on the plans", r"Steel door and frame"),
       C("OT-10441613-2100", 1, "Portable fire extinguisher, ABC all purpose, 20 lb. (Ea.)", "substitute", "extinguisher size and rating not specified on the plans", r"20 lb"),
       C("OT-10441613-9400", 1, "Installation of extinguishers, 12 or more, on nailable surface (Ea.)", "substitute", "row condition is 12 or more units; 8 units here", r"12 or more"))
unpriced("G23", "not_in_book", "the only verified grab-bar row (install only) carries no unit and is not usable; material rows not verified", "OT-10281313-9996 (unit blank)")
unpriced("G24", "not_in_book", "no verified toilet tissue dispenser row")
unpriced("G25", "not_in_book", "verified row OT-10281313-6050 carries unit S.F. (misread) and is not usable", "OT-10281313-6050 (unit unreadable)")
unpriced("G26", "not_in_book", "verified row OT-10281313-8100 carries no unit and is not usable", "OT-10281313-8100 (unit blank)")
priced("G27", C("OT-10281313-2300", 1, "Hand dryer, surface mounted, electric, 115 volt, 20 amp (Ea.)", "substitute", "dryer model and voltage not specified (A-5.12 KN11: 'GC to coordinate power supply')", r"Hand dryer, surface mounted"))
priced("G28", C("OT-10281313-4400", 1, "Robe hook, single, heavy duty, concealed mounting (Ea.)", "substitute", "coat hook product not specified (A-5.12 KN13)", r"Heavy duty, concealed mounting"))
priced("G29", C("OT-10281313-3200", 1, "Mirror with stainless steel 3/4\" square frame, 48\" x 24\" (Ea.)", "substitute", "mirror size and frame per A-7.20 not issued", r"48\" x 24"))
priced("G30", C("OT-10142313-2070", 1, "Engraved panel signage, interior, corridor, stock acrylic, 2-sided, 3\" x 12\" (Ea.)", "substitute",
                "drawn: toilet room identification signs per 8 & 9 / A-5.10 (sheet not issued); row is a stock 2-sided corridor sign, not a tactile ADA sign", r"3\" x 12"))
unpriced("G31", "not_quantified", "sign schedule not issued")
unpriced("G32", "not_quantified", "penetrations per M/P/E plans not issued")
unpriced("G33", "not_quantified", "not quantifiable")
unpriced("G34", "not_in_book", "survey service, no row")
unpriced("G35", "not_quantified", "not quantifiable")
unpriced("G36", "none_required", "no partitions drawn")

# ---------------------------------------------------------------------------------------------
# Join
# ---------------------------------------------------------------------------------------------
UNIT_ALIAS = {"SF": "S.F.", "LF": "L.F.", "EA": "Ea.", "RISER": "Riser", "WK": "Week", "DAY": "Day",
              "LS": "L.S.", "GAL": "Gal.", "MSF": "M.S.F.", "OPNG": "Opng."}
PRICED = ("priced_exact", "priced_substitute")


def unit_check(takeoff_unit, book_unit, factor_text):
    tu = UNIT_ALIAS.get(takeoff_unit, takeoff_unit)
    if factor_text == "1" and tu.lower().rstrip(".") == book_unit.lower().rstrip("."):
        return "same unit"
    if factor_text == "1" and {tu, book_unit} <= {"Ea.", "Door", "Opng."}:
        return "each-type units (takeoff Ea. = book " + book_unit + ")"
    return f"converted: book qty = takeoff qty x {factor_text}"


def empty_book():
    return dict(cost_code="", ref_line="", book_page="", book_description_path="", book_unit="", unit_check="", factor="", book_qty="",
                latest_period="", material_la_latest="", labor_la_latest="", equipment_la_latest="", total_la_latest="",
                total_incl_op_la_latest="", ext_material="", ext_labor="", ext_equipment="", ext_total="", ext_total_incl_op="")


EST_COLS = ["item_id", "bid_letter_item", "division", "description", "component", "qty", "unit", "qty_basis", "qty_source",
            "cost_code", "ref_line", "book_page", "book_description_path", "book_unit", "unit_check", "factor", "book_qty",
            "latest_period", "material_la_latest", "labor_la_latest", "equipment_la_latest", "total_la_latest", "total_incl_op_la_latest",
            "ext_material", "ext_labor", "ext_equipment", "ext_total", "ext_total_incl_op",
            "price_status", "match", "candidate_row", "reason", "takeoff_notes"]
rows = []
unpriced_rows = []        # fully unpriced lines and unpriced components
missing = [r["item_id"] for r in TO.ROWS if r["item_id"] not in P]
assert not missing, f"takeoff lines without a pricing decision: {missing}"
extra = [k for k in P if k not in {r["item_id"] for r in TO.ROWS}]
assert not extra, f"pricing decisions without a takeoff line: {extra}"
dup_notes = {}

for t in TO.ROWS:
    pm = P[t["item_id"]]
    base = dict(item_id=t["item_id"], bid_letter_item=t["bid_letter_item"], division=t["division"], description=t["description"],
                qty=t["qty"], unit=t["unit"], qty_basis=t["qty_basis"], qty_source=t["qty_source"], takeoff_notes=t["notes"])
    if pm["kind"] == "unpriced" or t["qty"] == "":
        status = pm["status"] if pm["kind"] == "unpriced" else "not_quantified"
        reason = pm.get("reason", "quantity not determinable") if pm["kind"] == "unpriced" else "quantity not determinable"
        r = dict(base, component="", **empty_book(), price_status=status, match="", candidate_row=pm.get("candidate", ""), reason=reason)
        rows.append(r)
        if status != "none_required":
            unpriced_rows.append(dict(r, scope="line"))
        continue
    qty = Decimal(str(t["qty"]))
    if all(c["kind"] == "ucomp" for c in pm["comps"]):
        # a line whose only component has no usable row is unpriced as a whole (scope = line)
        assert len(pm["comps"]) == 1, f"{t['item_id']}: several unpriced components and no priced one; split the line or merge the reasons"
        c = pm["comps"][0]
        r = dict(base, component="", **empty_book(), price_status=c["status"], match="", candidate_row=c["candidate"], reason=c["desc"] + ": " + c["reason"])
        rows.append(r)
        unpriced_rows.append(dict(r, scope="line"))
        continue
    for c in pm["comps"]:
        if c["kind"] == "ucomp":
            r = dict(base, component=c["desc"], **empty_book(), price_status=c["status"], match="", candidate_row=c["candidate"], reason=c["reason"])
            rows.append(r)
            unpriced_rows.append(dict(r, scope="component"))
            continue
        cc = c["cc"]
        allrows = BOOKROWS.get(cc, [])
        # a code that is absent from the table is a hard stop: the component must be reclassified by hand
        # (unpriced with a candidate_row), never priced from memory of an earlier table
        assert allrows, f"{t['item_id']}: cost code {cc} is not carried by any row of the table on hand; reclassify the component"
        ver = [b for b in allrows if b["status"] == "verified"]
        dup = ""
        if len(allrows) > 1:
            dup = f"cost code appears on {len(allrows)} rows of the table: " + "; ".join(
                f"[{b['status']}, unit '{b['unit'] or '-'}'] {b['description'][:60]}" for b in allrows)
            dup_notes[cc] = dup
        if len(ver) != 1:
            r = dict(base, component=c["desc"], **empty_book(), price_status="cost_code_duplicated", match="",
                     candidate_row=cc, reason="cost code duplicated in book: " + (dup or "no single verified row") + "; " + c["reason"])
            rows.append(r)
            unpriced_rows.append(dict(r, scope="component"))
            continue
        b = ver[0]
        assert re.search(c["expect"], b["description"]), f"{t['item_id']}: {cc} verified row description '{b['description']}' does not match /{c['expect']}/"
        assert b["total_la_latest"] != "", f"{t['item_id']}: {cc} has no total_la_latest"
        assert b["unit"] != "", f"{t['item_id']}: {cc} has no unit"
        f_ = c["factor"] if isinstance(c["factor"], Decimal) else Decimal(str(c["factor"]))
        factor_text = str(c["factor"]) if not isinstance(c["factor"], Decimal) else f"{f_:.6f}"
        bq = qty * f_
        m, l, e, tot, op = (D(b["material_la_latest"]) or Decimal(0), D(b["labor_la_latest"]) or Decimal(0),
                            D(b["equipment_la_latest"]) or Decimal(0), D(b["total_la_latest"]), D(b["total_incl_op_la_latest"]))
        assert q2(m + l + e) == q2(tot), f"{cc}: components {m}+{l}+{e} do not sum to total {tot}"
        r = dict(base, component=c["desc"], cost_code=cc, ref_line=b["ref_line"], book_page=b["book_page"], book_description_path=b["description_path"],
                 book_unit=b["unit"], unit_check=unit_check(t["unit"], b["unit"], factor_text), factor=factor_text,
                 # the book quantity is written with enough digits that book_qty x unit price reproduces the extended amount
                 book_qty=(str(q2(bq)) if q2(bq) == bq else f"{bq:.6f}"),
                 latest_period=b["latest_period"], material_la_latest=str(m), labor_la_latest=str(l), equipment_la_latest=str(e),
                 total_la_latest=str(tot), total_incl_op_la_latest=(str(op) if op is not None else ""),
                 ext_material=str(q2(bq * m)), ext_labor=str(q2(bq * l)), ext_equipment=str(q2(bq * e)), ext_total=str(q2(bq * tot)),
                 ext_total_incl_op=(str(q2(bq * op)) if op is not None else ""),
                 price_status=("priced_exact" if c["tier"] == "exact" else "priced_substitute"), match=c["match"], candidate_row="",
                 reason=c["reason"] + ("; " + dup if dup else ""))
        rows.append(r)

with open(OUT_CSV, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=EST_COLS)
    w.writeheader()
    for r in rows:
        w.writerow(r)

UNP_COLS = ["item_id", "scope", "bid_letter_item", "division", "description", "component", "qty", "unit", "qty_basis", "qty_source",
            "price_status", "reason", "candidate_row", "takeoff_notes"]
with open(OUT_UNP, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=UNP_COLS, extrasaction="ignore")
    w.writeheader()
    for r in unpriced_rows:
        w.writerow(r)

# ---------------------------------------------------------------------------------------------
# Totals by tier (direct = sum of bare ext_total; company markup applied once, on the bare direct total of each tier)
# ---------------------------------------------------------------------------------------------
CONT, OH, PROF = Decimal("0.04"), Decimal("0.08"), Decimal("0.07")
SELL_FACTOR = (1 + CONT) * (1 + OH) * (1 + PROF)          # 1.04 x 1.08 x 1.07 = 1.2018240
ITEMS = ["Demolition", "Interior walls", "Ceilings", "Flooring", "New Roofing", "Interior finishes", "Exterior finishes",
         "Doors & Hardware", "Stairs & railing", "Owner-provided equipment install", "Lighting", "Mechanical", "Plumbing",
         "Low Voltage", "F&B equipment", "Other required scope"]
TIERS = ("exact", "substitute")


def zero():
    return dict(material=Decimal(0), labor=Decimal(0), equipment=Decimal(0), direct=Decimal(0), book_op=Decimal(0), rows_without_op=0, components=0)


by_item = {k: dict(exact=zero(), substitute=zero(), lines=set(), lines_with_price=set(), lines_unpriced=set(), unpriced_components=0) for k in ITEMS}
for r in rows:
    g = by_item[r["bid_letter_item"]]
    g["lines"].add(r["item_id"])
    if r["price_status"] in PRICED:
        tier = "exact" if r["price_status"] == "priced_exact" else "substitute"
        s = g[tier]
        s["components"] += 1
        g["lines_with_price"].add(r["item_id"])
        s["material"] += Decimal(r["ext_material"]); s["labor"] += Decimal(r["ext_labor"]); s["equipment"] += Decimal(r["ext_equipment"])
        s["direct"] += Decimal(r["ext_total"])
        if r["ext_total_incl_op"] != "":
            s["book_op"] += Decimal(r["ext_total_incl_op"])
        else:
            s["rows_without_op"] += 1
    elif r["price_status"] != "none_required":
        if r["component"]:
            g["unpriced_components"] += 1
        else:
            g["lines_unpriced"].add(r["item_id"])


def tier_total(tier):
    return dict(material=str(sum(g[tier]["material"] for g in by_item.values())), labor=str(sum(g[tier]["labor"] for g in by_item.values())),
                equipment=str(sum(g[tier]["equipment"] for g in by_item.values())), direct=str(sum(g[tier]["direct"] for g in by_item.values())),
                book_op=str(sum(g[tier]["book_op"] for g in by_item.values())), rows_without_op=sum(g[tier]["rows_without_op"] for g in by_item.values()),
                components=sum(g[tier]["components"] for g in by_item.values()),
                sell=str(q2(sum(g[tier]["direct"] for g in by_item.values()) * SELL_FACTOR)))


tot = dict(exact=tier_total("exact"), substitute=tier_total("substitute"))
direct_all = Decimal(tot["exact"]["direct"]) + Decimal(tot["substitute"]["direct"])
tot["combined"] = dict(direct=str(direct_all), sell=str(q2(direct_all * SELL_FACTOR)),
                       book_op=str(Decimal(tot["exact"]["book_op"]) + Decimal(tot["substitute"]["book_op"])),
                       contingency=str(q2(direct_all * CONT)), overhead=str(q2(direct_all * (1 + CONT) * OH)),
                       profit=str(q2(direct_all * (1 + CONT) * (1 + OH) * PROF)))
tot["sell_factor"] = str(SELL_FACTOR)
tot["items"] = {k: dict(exact={kk: (str(vv) if isinstance(vv, Decimal) else vv) for kk, vv in g["exact"].items()},
                        substitute={kk: (str(vv) if isinstance(vv, Decimal) else vv) for kk, vv in g["substitute"].items()},
                        lines=sorted(g["lines"]), lines_with_price=sorted(g["lines_with_price"]), lines_unpriced=sorted(g["lines_unpriced"]),
                        unpriced_components=g["unpriced_components"]) for k, g in by_item.items()}
tot["n_takeoff_lines"] = len(TO.ROWS)
tot["n_estimate_rows"] = len(rows)
tot["n_unpriced_lines"] = sum(1 for r in unpriced_rows if r["scope"] == "line")
tot["n_unpriced_components"] = sum(1 for r in unpriced_rows if r["scope"] == "component")
tot["duplicate_cost_codes_seen"] = dup_notes
# Components whose classification changed only because the book table was replaced (one cost code per printed
# line; lines under an unreadable section header carry no code). Recorded for the summary; no price effect beyond
# what the rows above already show.
tot["changed_by_table_replacement"] = [
    dict(item_id="I02", component="poplar crown molding", before="priced_substitute (OT-06221330-2510, verified in the previous table)",
         after="book_row_not_verified (p203 line 2510 is flagged and carries no cost code)"),
    dict(item_id="G28", component="robe hook OT-10281313-4400", before="cost_code_duplicated (code shared with a flagged towel-bar row)",
         after="priced_substitute (the code is now carried by the verified robe-hook row only)"),
    dict(item_id="P03", component="grease interceptor (candidate only)", before="candidate OT-22131913-0150 / -0200 / -1060 / -1120 (verified)",
         after="candidate p514 lines 0150 / 0200 / 1060 / 1120 (flagged, no cost code); line stays spec_not_issued"),
    dict(item_id="D10", component="aluminum window removal (candidate only)", before="candidate OT-08050510-0200 (machine_read window row sharing the code with a flagged door row)",
         after="candidate p268 window line 0200 (flagged, no cost code); line stays book_row_not_verified"),
]
json.dump(tot, open(OUT_TOT, "w"), indent=1)

# ---------------------------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------------------------
from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook()
ws = wb.active
ws.title = "estimate_la_latest"
ws.append(EST_COLS)
NUM = {"qty", "book_qty", "material_la_latest", "labor_la_latest", "equipment_la_latest", "total_la_latest", "total_incl_op_la_latest",
       "ext_material", "ext_labor", "ext_equipment", "ext_total", "ext_total_incl_op"}
for r in rows:
    ws.append([(float(r[c]) if (c in NUM and r[c] not in ("", None)) else r[c]) for c in EST_COLS])
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "E2"
ws2 = wb.create_sheet("takeoff")
tcols = ["item_id", "bid_letter_item", "division", "description", "qty", "unit", "qty_basis", "qty_source", "notes"]
ws2.append(tcols)
for t in TO.ROWS:
    ws2.append([(float(t[c]) if (c == "qty" and t[c] != "") else t[c]) for c in tcols])
for c in ws2[1]:
    c.font = Font(bold=True)
ws3 = wb.create_sheet("unpriced_items")
ws3.append(UNP_COLS)
for r in unpriced_rows:
    ws3.append([(float(r[c]) if (c == "qty" and r[c] != "") else r[c]) for c in UNP_COLS])
for c in ws3[1]:
    c.font = Font(bold=True)
ws4 = wb.create_sheet("summary_by_item")
ws4.append(["bid_letter_item", "exact direct (bare)", "exact sell", "substitute direct (bare)", "substitute sell", "combined direct", "combined sell",
            "book O&P extended (reference)", "lines with a price", "unpriced lines", "unpriced components"])
for k in ITEMS:
    g = by_item[k]
    ex, su = g["exact"]["direct"], g["substitute"]["direct"]
    ws4.append([k, float(ex), float(q2(ex * SELL_FACTOR)), float(su), float(q2(su * SELL_FACTOR)), float(ex + su), float(q2((ex + su) * SELL_FACTOR)),
                float(g["exact"]["book_op"] + g["substitute"]["book_op"]), len(g["lines_with_price"]), len(g["lines_unpriced"]), g["unpriced_components"]])
ws4.append(["TOTAL", float(tot["exact"]["direct"]), float(tot["exact"]["sell"]), float(tot["substitute"]["direct"]), float(tot["substitute"]["sell"]),
            float(tot["combined"]["direct"]), float(tot["combined"]["sell"]), float(tot["combined"]["book_op"]), "", tot["n_unpriced_lines"], tot["n_unpriced_components"]])
ws4.append([])
ws4.append(["priced_exact: the only tier priced from the book without an assumption (row = specified work and product as the issued sheets fix it)."])
ws4.append(["priced_substitute: a different or unspecified product, size, gauge, type or finish system, or part coverage; for decision only (see reason column)."])
ws4.append(["Sell = direct x 1.04 (contingency) x 1.08 (overhead) x 1.07 (profit) = direct x " + str(SELL_FACTOR) + "; applied once on the bare direct of each tier."])
ws4.append(["Book O&P column is the installing contractor's O&P; shown for reference only, not added to any sell figure."])
for c in ws4[1]:
    c.font = Font(bold=True)
for w_ in (ws, ws2, ws3, ws4):
    for col in w_.columns:
        w_.column_dimensions[col[0].column_letter].width = 18
wb.save(OUT_XLSX)

print(f"estimate rows {len(rows)}: exact components {tot['exact']['components']}, substitute components {tot['substitute']['components']}; "
      f"unpriced lines {tot['n_unpriced_lines']}, unpriced components {tot['n_unpriced_components']}; duplicate codes seen {len(dup_notes)}")
print(f"exact direct {tot['exact']['direct']} sell {tot['exact']['sell']} | substitute direct {tot['substitute']['direct']} sell {tot['substitute']['sell']} | combined direct {tot['combined']['direct']} sell {tot['combined']['sell']}")
for cc, d in dup_notes.items():
    print("  duplicate:", cc, d[:160])
print("written:", OUT_CSV, OUT_XLSX, OUT_UNP, OUT_TOT)

"""Writes rfi_list.csv and rfi_package_DRAFT.xlsx from rfi_data.py.
Usage (from takeoff2/): python3 -I scripts/build_rfi_package.py
The xlsx is a copy of RFI_Template_KellyArchitects.xlsx with one sheet per RFI; header cells
filled, question filled, answer cells left blank, no prices. DRAFT - nothing is sent.
"""
import csv, io, os, sys, datetime
from copy import copy
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor

HERE = os.path.dirname(os.path.abspath(__file__))
T2 = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from rfi_data import RFIS  # noqa: E402

DATE_CREATED = datetime.date(2026, 10, 28)
AUTHOR_CO = "Owl Tower Construction Group, Inc."
AUTHOR_ADDR1 = "CSLB #1108551"
AUTHOR_ADDR2 = "Torrance, CA"
AUTHOR_BY = "Estimating, Owl Tower Construction Group, Inc."

# ---------------- CSV ----------------
csv_path = os.path.join(T2, "rfi_list.csv")
with open(csv_path, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["rfi_no", "sheets_documents", "discipline", "subject", "question", "why_it_matters",
                "importance", "source_note_file"])
    for r in RFIS:
        w.writerow([f"{r['no']:03d}", r["sheets"], r["discipline"], r["subject"], r["question"],
                    r["why"], r["importance"], "; ".join(r["sources"])])
print("wrote", csv_path, len(RFIS), "rows")

# ---------------- XLSX ----------------
tpl_path = os.path.join(T2, "RFI_Template_KellyArchitects.xlsx")
wb = openpyxl.load_workbook(tpl_path)
tpl = wb["RFI #000"]
logo = tpl._images[0]
logo_bytes = logo._data()
logo_from = (logo.anchor._from.col, logo.anchor._from.row)
logo_to = (logo.anchor.to.col, logo.anchor.to.row)
tpl_print_area = tpl.print_area
tpl_setup = copy(tpl.page_setup)
tpl_margins = copy(tpl.page_margins)
tpl_props = copy(tpl.sheet_properties)
tpl_print_opts = copy(tpl.print_options)


def fill(ws, r):
    ws.title = f"RFI #{r['no']:03d}"
    ws["A2"] = f"Request for Information #{r['no']:03d}"
    ws["F5"] = r["importance"]                       # Importance: High/Med/Low
    ws["J5"] = DATE_CREATED                           # Date Created (number format mm-dd-yy kept)
    ws["A16"] = AUTHOR_CO                             # Author Company
    ws["A17"] = AUTHOR_ADDR1
    ws["A18"] = AUTHOR_ADDR2
    ws["A19"] = "Phone:"
    ws["E16"] = AUTHOR_BY                             # Authored By
    ws["A21"] = r["subject"]                          # Subject
    ws["E21"] = r["discipline"]                       # Discipline
    ws["I21"] = r["category"]                         # Category
    ws["A30"] = r["question"] + f"\n\nSheets / documents: {r['sheets']}.\nBid impact: {r['why']}"
    ws["A37"] = r["suggestion"] if r["suggestion"] else None
    ws["A42"] = None                                  # Answer left blank
    ws["J29"] = None                                  # Date Required left blank
    ws["J41"] = None                                  # Date Answered left blank
    # question block: wrap and top-align, make the merged block tall enough
    ws["A30"].alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top", horizontal="left")
    n_lines = max(6, len(r["question"]) // 95 + len(r["why"]) // 95 + 5)
    for row in range(30, 36):
        ws.row_dimensions[row].height = max(12.0, n_lines * 12.75 / 6)
    # logo, print area and page setup are not carried by copy_worksheet
    img = XLImage(io.BytesIO(logo_bytes))
    img.width, img.height = logo.width, logo.height
    anc = TwoCellAnchor()
    anc._from.col, anc._from.row = logo_from
    anc.to.col, anc.to.row = logo_to
    img.anchor = anc
    ws.add_image(img)
    ws.print_area = tpl_print_area.split("!")[-1] if tpl_print_area else "A1:J46"
    ws.page_setup = copy(tpl_setup)
    ws.page_margins = copy(tpl_margins)
    ws.sheet_properties = copy(tpl_props)
    ws.print_options = copy(tpl_print_opts)


for r in RFIS:
    ws = wb.copy_worksheet(tpl)
    fill(ws, r)

wb.remove(tpl)
out = os.path.join(T2, "rfi_package_DRAFT.xlsx")
wb.save(out)
print("wrote", out, "sheets:", len(wb.sheetnames))

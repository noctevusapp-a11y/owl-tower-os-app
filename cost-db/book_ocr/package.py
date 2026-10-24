#!/usr/bin/env python3
"""Copy the finished book table into the repository data folder and write the XLSX for the office.

  python3 -I package.py <table_dir> <spot_dir> <dest_dir>
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def main() -> int:
    table, spot, dest = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    dest.mkdir(parents=True, exist_ok=True)
    for name in ("book_lines_la_latest.csv", "book_lines_la_latest.json", "book_sections.csv", "build_report.json", "derive_report.json"):
        shutil.copy(table / name, dest / name)
    (dest / "spot_check").mkdir(exist_ok=True)
    for p in spot.glob("*"):
        shutil.copy(p, dest / "spot_check" / p.name)
    rows = list(csv.DictReader((table / "book_lines_la_latest.csv").open(encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = "lines"
    cols = ["cost_code", "division", "division_name", "section_id", "section_title", "line", "description", "description_path", "unit", "crew",
            "daily_output", "labor_hours", "material_2022", "labor_2022", "equipment_2022", "total_2022", "total_incl_op_2022",
            "material_add_pct", "labor_add_pct", "cci_row", "material_factor_id", "material_factor", "labor_factor", "equipment_factor",
            "material_la_latest", "labor_la_latest", "equipment_la_latest", "total_la_latest", "total_incl_op_la_latest", "latest_period",
            "status", "flags", "book_page", "ref_line"]
    ws.append(cols)
    num = {"daily_output", "labor_hours", "material_2022", "labor_2022", "equipment_2022", "total_2022", "total_incl_op_2022", "material_factor",
           "labor_factor", "equipment_factor", "material_la_latest", "labor_la_latest", "equipment_la_latest", "total_la_latest", "total_incl_op_la_latest",
           "book_page", "material_add_pct", "labor_add_pct"}
    for r in rows:
        out = []
        for c in cols:
            v = r.get(c)
            if v in (None, ""):
                out.append(None)
            elif c in num:
                try:
                    out.append(float(v))
                except ValueError:
                    out.append(v)
            else:
                out.append(v)
        ws.append(out)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    widths = {"description": 50, "description_path": 70, "section_title": 30, "flags": 30, "cost_code": 22, "division_name": 26}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    ws2 = wb.create_sheet("sections")
    with (table / "book_sections.csv").open(encoding="utf-8") as f:
        for row in csv.reader(f):
            ws2.append(row)
    ws3 = wb.create_sheet("readme")
    for line in ("Owl Tower book table: 2022 national bare costs read from the company's reference copy of the 2022 cost book, with derived Los Angeles latest-month estimates.",
                 "status = verified: every number supported by two independent readings or the table arithmetic; flagged: check against the book page before use.",
                 "*_la_latest = component x City Cost Index (LA, 2022) x BLS factor per component; derived, unreviewed. Not quotes, not observed 2026 prices.",
                 "Column dictionary: COLUMNS.md. Report: extraction_report.md."):
        ws3.append([line])
    wb.save(dest / "book_lines_la_latest.xlsx")
    rep = json.loads((table / "build_report.json").read_text(encoding="utf-8"))
    print(f"packaged {len(rows)} rows; verified {rep.get('verified')} flagged {rep.get('flagged')} -> {dest}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
